# -*- coding: utf-8 -*-
r"""용인시 데모 2차 DB — 정본 스냅숏(base) + 관리층(mgmt) 을 판으로 나눠 쌓는다. 정본은 읽기만 한다.

    python build_demo_db.py base          정본 발행 판을 base_<릴리스>\ 로 복사 (이미 있으면 멈춘다)
    python build_demo_db.py mgmt v1.0     관리층 mgmt_v1.0_<날짜>\ 생성 (이미 있으면 멈춘다)
    python build_demo_db.py check         모든 판의 _MANIFEST 지문을 다시 재어 바뀐 파일이 없는지 확인

왜 이렇게 나누나 (사용자 2026-09-20)
  「의무DB뿐만 아니라 관리를 위한 데이터(관리대상DB)나 관리용 필드도 추가될 거라 버전을 달리 관리하자.」
  · base  = 1차 데모(2026-09-05 개발 인계, `..\_데모_용인시_20260905\`)와 **같은 맥락·같은 폴더 구성**의 정본 스냅숏.
            정본 릴리스 하나에 폴더 하나. 한 번 만들면 고치지 않는다.
  · mgmt  = 우리가 얹는 관리 데이터·관리 필드(3단 분류·관리영역·필터·관리대상…). 바뀔 때마다 새 판 폴더.
            정본 표의 열을 고치지 않고 obl_id 등 정본 키로 붙는 **별도 표**로 둔다 → 정본이 바뀌어도 조인만 다시 하면 된다.
  · 판을 덮어쓰지 않는다. 옛 판은 그대로 두고 _CURRENT.json 이 쓰는 판을 가리킨다.

1차 패키지와 다른 점
  · 「최신 파일」을 수정시각·파일명으로 고르지 않고 정본 `build\_RELEASE.json` 발행 판만 읽는다(E66).
  · 03 폴더 이름을 「관리대상」으로(용어 규칙 2026-09-11). 파일명은 정본 그대로.
  · 정본에 새로 생긴 표 doc_authority · target_type_all · target_catalog 를 함께 담는다.
"""
import csv
import datetime
import hashlib
import io
import json
import os
import shutil
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
csv.field_size_limit(10 ** 9)

HERE = os.path.dirname(os.path.abspath(__file__))
DEMO = os.path.dirname(HERE)                                   # ...\_데모_용인시_20260920
ROOT = os.path.dirname(DEMO)                                   # ...\ADOMS_DB_v1
BUILD = os.path.join(ROOT, "build")
OLD1 = os.path.join(ROOT, "_데모_용인시_20260905")
WORK = os.path.join(os.path.dirname(ROOT), "_체크리스트_시설별")   # 3단 분류 작업 폴더
TODAY = datetime.date.today().strftime("%Y%m%d")

BASE_PICK = [
    ("01_사실층", ["law", "doc", "doc_authority", "unit", "schedule"]),
    ("02_판단층", ["obl_master", "obl_item", "obl_evidence", "unit_role", "obl_axes", "appl_rule"]),
    ("03_관리대상", ["target_type", "target_condition", "target_note", "target_xlaw", "target_type_all", "target_catalog"]),
    ("04_관계", ["edge", "mutatis", "mutatis_subst"]),
]


def sha(p):
    h = hashlib.sha1()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()[:12]


def rows_of(p):
    if not p.endswith(".csv"):
        return ""
    with io.open(p, encoding="utf-8-sig", errors="replace", newline="") as f:
        return max(sum(1 for _ in csv.reader(f)) - 1, 0)


def write_manifest(folder, man):
    with io.open(os.path.join(folder, "_MANIFEST.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["구분", "표", "파일", "행수", "sha1", "원본"])
        w.writeheader()
        for r in man:
            w.writerow(r)


def load_current():
    p = os.path.join(DEMO, "_CURRENT.json")
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else {}


def save_current(cur):
    cur["updated_at"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    json.dump(cur, open(os.path.join(DEMO, "_CURRENT.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)


# ── base ─────────────────────────────────────────────────────────
def build_base():
    rel = json.load(open(os.path.join(BUILD, "_RELEASE.json"), encoding="utf-8"))
    if rel.get("state") != "stable":
        sys.exit("정본 발행 판이 stable 이 아니다 — 멈춤 (%s)" % rel.get("state"))
    tag = rel["release"]
    out = os.path.join(DEMO, "base_%s" % tag)
    if os.path.exists(out):
        sys.exit("이미 있다 — 덮어쓰지 않는다: %s" % out)
    man = []
    for sub, names in BASE_PICK:
        os.makedirs(os.path.join(out, sub))
        for nm in names:
            t = rel["tables"].get(nm)
            if not t:
                print("   %-18s ★ 정본 발행 판에 없음" % nm)
                continue
            src = os.path.join(BUILD, t["file"])
            dst = os.path.join(out, sub, t["file"])
            shutil.copy2(src, dst)
            n = rows_of(dst)
            assert n == t["rows"], (nm, n, t["rows"])          # 발행 판에 적힌 행수와 같아야 한다
            man.append({"구분": sub, "표": nm, "파일": t["file"], "행수": n, "sha1": sha(dst),
                        "원본": "build\\" + t["file"]})
            print("   %-10s %-18s %-44s %s행" % (sub, nm, t["file"][:44], "{:,}".format(n)))
    # 05 — 1차 데모 전용 산출물(검수 전) 승계
    sub = "05_데모전용"
    os.makedirs(os.path.join(out, sub))
    src_dir = os.path.join(OLD1, "05_데모전용")
    for f in sorted(os.listdir(src_dir)):
        src = os.path.join(src_dir, f)
        dst = os.path.join(out, sub, f)
        shutil.copy2(src, dst)
        man.append({"구분": sub, "표": "1차 데모 승계(검수 전)", "파일": f, "행수": rows_of(dst), "sha1": sha(dst),
                    "원본": "_데모_용인시_20260905\\05_데모전용\\" + f})
        print("   %-10s %-18s %s" % (sub, "1차 승계", f[:44]))
    write_manifest(out, man)
    io.open(os.path.join(out, "_README.md"), "w", encoding="utf-8").write(BASE_README.format(
        tag=tag, pub=rel.get("published_at", ""), now=datetime.datetime.now().strftime("%Y-%m-%d %H:%M")))
    cur = load_current()
    cur["base"] = os.path.basename(out)
    cur["canon_release"] = tag
    save_current(cur)
    print("■ base 완료 →", out)


BASE_README = """# base_{tag} — 정본 스냅숏 (용인시 데모 2차)

**정본이 아니다.** 정본 발행 판 `{tag}`(발행 {pub})을 그대로 복사한 스냅숏이다. **이 폴더는 고치지 않는다.**
정본이 새로 발행되면 `base_<새 릴리스>\\` 를 새로 만든다.

폴더 구성은 1차 데모 인계 DB(`..\\..\\_데모_용인시_20260905\\`)와 같다. 차이: 03 이름이 「관리대상」 ·
새 표 `doc_authority`(문서-위임근거 다중) · `target_type_all`(관리대상 전 법령) · `target_catalog` 추가 · 05 는 1차 산출물 승계.
파일명은 정본 그대로 — 정본과 대조할 수 있게. 행수·sha1 은 `_MANIFEST.csv`.

생성 {now} · `_build\\build_demo_db.py base`
"""


# ── mgmt ─────────────────────────────────────────────────────────
def build_mgmt(ver):
    cur = load_current()
    if "base" not in cur:
        sys.exit("base 를 먼저 만든다")
    out = os.path.join(DEMO, "mgmt_%s_%s" % (ver, TODAY))
    if any(d.startswith("mgmt_%s_" % ver) for d in os.listdir(DEMO)):
        sys.exit("이 판(%s)은 이미 있다 — 새 판 번호를 쓴다" % ver)
    os.makedirs(out)
    sys.path.insert(0, WORK)
    from openpyxl import load_workbook
    from duty_master import MASTER, AREA_NAME, FLAG, VERSION
    from law_family import FAMILY, family_of

    base = os.path.join(DEMO, cur["base"])
    rd = lambda sub, name: list(csv.DictReader(io.open(os.path.join(base, sub, name), encoding="utf-8-sig", newline="")))
    rel = json.load(open(os.path.join(BUILD, "_RELEASE.json"), encoding="utf-8"))
    om = {o["obl_id"]: o for o in rd("02_판단층", rel["tables"]["obl_master"]["file"])}
    laws = rd("01_사실층", rel["tables"]["law"]["file"])

    man = []

    def dump(name, head, rows, note):
        p = os.path.join(out, name)
        with io.open(p, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(head)
            w.writerows(rows)
        man.append({"구분": "관리층", "표": note, "파일": name, "행수": len(rows), "sha1": sha(p), "원본": "생성"})
        print("   %-44s %s행" % (name, "{:,}".format(len(rows))))

    # ① 중처법 의무조항 마스터 36
    dump("duty_article.csv", ["area", "area_name", "code", "name_ko", "sapa_basis", "checklist_no", "flag"],
         [[m[0], AREA_NAME[m[0]], m[1], m[2], m[3], m[4], FLAG.get(m[1], "")] for m in MASTER],
         "중처법 의무조항 36 (표준 명칭 %s)" % VERSION)

    # ② 의무 ↔ 1단 카테고리 (영역별 중복 허용 · 한 행 = 의무 × 영역)
    wb = load_workbook(os.path.join(WORK, "(분석)체크리스트_정본대조_20260920.xlsx"), read_only=True)
    rows = []
    for area, sheet in (("I", "산업재해"), ("F", "시민_시설교통"), ("M", "시민_원료제조물")):
        it = wb[sheet].iter_rows(values_only=True)
        h = next(it)
        ix = {k: h.index(k) for k in h if k}
        for r in it:
            oid = r[ix["3단 obl_id"]]
            if not oid:
                continue
            abbr, title = r[ix["법(약칭)"]] or "", r[ix["법(정식 명칭)"]] or ""
            fam = family_of(abbr, title)
            rows.append([oid, area, r[ix["1단 코드"]], r[ix["근거유형"]], r[ix["규칙·정답"]] or "",
                         r[ix["2단 키(문서|조|가지)"]] or "", r[ix["2단 세부 이행업무"]] or "",
                         r[ix["영역 근거"]] or "", r[ix["확신도"]] or "", fam, FAMILY[fam],
                         om.get(oid, {}).get("obligation_group", ""), "Y" if r[ix["행정청 권한"]] else "",
                         "confirmed" if r[ix["근거유형"]] in ("중처법조문", "체크리스트") else "pending"])
    dump("obl_category.csv",
         ["obl_id", "area", "code", "basis_type", "basis_detail", "task_key", "task_name", "area_basis",
          "area_confidence", "law_family", "law_family_name", "mg", "admin_power", "review_status"],
         rows, "의무 × 관리영역 → 1단 카테고리 · 2단 업무 · 법령 가족(필터 A) · MG")

    # ③ 영역 미배정 (삭제 아님)
    it = wb["영역미배정"].iter_rows(values_only=True)
    h = next(it)
    ix = {k: h.index(k) for k in h if k}
    dump("obl_unassigned.csv", ["obl_id", "reason", "law_abbr", "doc", "unit_label"],
         [[r[ix["obl_id"]], r[ix["사유"]], r[ix["법(약칭)"]], r[ix["문서"]], r[ix["조문"]]] for r in it if r[ix["obl_id"]]],
         "어느 영역에도 들지 않은 의무 — 검수 대상(삭제 아님)")

    # ④ 체크리스트 정답(교육 자료 시설별 10종) 3단 매핑
    it = wb["3단매핑_시범"].iter_rows(values_only=True)
    h = next(it)
    dump("checklist_gold.csv", list(h), [list(r) for r in it], "교육 자료 체크리스트 195 → 1단·2단·3단(시범)")

    # ⑤ 법 → 관리영역 · 법령 가족
    it = wb["영역배정_법"].iter_rows(values_only=True)
    h = next(it)
    ix = {k: h.index(k) for k in h if k}
    area_rows = {r[ix["법(정식 명칭)"]]: r for r in it}
    lrows, miss = [], 0
    for l in laws:
        if l.get("status") == "merged":
            continue
        fam = family_of(l["abbr_ko"], l["title_ko"])
        a = area_rows.get(l["title_ko"])
        miss += a is None
        g = (lambda k: (a[ix[k]] or "") if a else "")
        lrows.append([l["law_id"], l["abbr_ko"], l["title_ko"], fam, FAMILY[fam],
                      "Y" if g("산업") else "", "Y" if g("시설·교통") else "", "Y" if g("원료·제조물") else "",
                      g("확신도"), g("근거")])
    assert miss == 0, "영역배정_법 에 없는 법 %d" % miss
    dump("law_mgmt.csv", ["law_id", "abbr_ko", "title_ko", "law_family", "law_family_name",
                          "area_I", "area_F", "area_M", "area_confidence", "area_basis"],
         lrows, "법 → 법령 가족 · 관리영역(검수 전)")

    # ⑥ 코드표
    dump("code_law_family.csv", ["code", "name_ko"], sorted(FAMILY.items()), "필터 A 법령 가족 15")

    write_manifest(out, man)
    io.open(os.path.join(out, "_README.md"), "w", encoding="utf-8").write(MGMT_README.format(
        ver=ver, base=cur["base"], std=VERSION, now=datetime.datetime.now().strftime("%Y-%m-%d %H:%M")))
    cur.setdefault("mgmt_history", []).append(os.path.basename(out))
    cur["mgmt"] = os.path.basename(out)
    save_current(cur)
    print("■ mgmt 완료 →", out)


MGMT_README = """# mgmt_{ver} — 관리층 (용인시 데모 2차)

정본 스냅숏 `{base}` 위에 얹는 **관리 데이터·관리 필드**. 정본 표를 고치지 않고 정본 키(`obl_id`·`law_id`)로 붙는 별도 표다.
**판단층이다 — 검수 전 값이 섞여 있다**(`review_status`). 이 판은 고치지 않는다. 바꾸면 새 판(`mgmt_v1.1_…`)을 만든다.

| 파일 | 내용 | 키 |
|---|---|---|
| `duty_article.csv` | 중처법 의무조항 36 — 표준 명칭 {std} · 원료·제조물 「도급·위탁 공정 포함」 표시 | `code` |
| `obl_category.csv` | 의무 × 관리영역 → 1단 카테고리 · 2단 세부 업무 · 법령 가족(필터 A) · 의무군 MG | `obl_id`+`area` |
| `obl_unassigned.csv` | 영역 미배정 의무 — 검수 대상(**삭제 아님**) | `obl_id` |
| `checklist_gold.csv` | 교육 자료 체크리스트(시설별 10종) 195 → 3단 매핑 시범(정답) | — |
| `law_mgmt.csv` | 법 → 법령 가족 · 관리영역 · 확신도 | `law_id` |
| `code_law_family.csv` | 법령 가족 코드표 15 | `code` |

`review_status` — `confirmed` = 중처법 조문 번호 또는 체크리스트 정답으로 확정 · `pending` = 낱말 규칙·기본값 제안(사람 검수 필요).

생성 {now} · `_build\\build_demo_db.py mgmt {ver}` · 원천 `30_데이터\\_수집작업\\_체크리스트_시설별\\`
"""


def check():
    bad = 0
    for d in sorted(os.listdir(DEMO)):
        mf = os.path.join(DEMO, d, "_MANIFEST.csv")
        if not os.path.exists(mf):
            continue
        for r in csv.DictReader(io.open(mf, encoding="utf-8-sig")):
            p = os.path.join(DEMO, d, r["구분"], r["파일"]) if r["구분"] != "관리층" else os.path.join(DEMO, d, r["파일"])
            ok = os.path.exists(p) and sha(p) == r["sha1"]
            bad += not ok
            if not ok:
                print("   ★ 바뀜/없음:", d, r["파일"])
        print("   %s 확인" % d)
    print("■ 이상 없음" if not bad else "■ %d 건 다름" % bad)


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "base":
        build_base()
    elif cmd == "mgmt" and len(sys.argv) > 2:
        build_mgmt(sys.argv[2])
    elif cmd == "check":
        check()
    else:
        print(__doc__)
