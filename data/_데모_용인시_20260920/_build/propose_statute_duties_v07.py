# -*- coding: utf-8 -*-
r"""데모 의무 자료 — 핵심 관계법령 법률·시행령·시행규칙 조문 의무 편입 **후보·표본 검수표** (2026-09-21 · v0.7 준비)

    python propose_statute_duties_v07.py            (결과 파일이 이미 있으면 멈춘다)
    python propose_statute_duties_v07.py --overwrite

왜
  데모 의무 자료(duty_class)에 재난안전법·건축물관리법·소방시설법·전기안전관리법·승강기법·철도안전법 등의
  **법률 조문 의무가 거의 없고 고시·훈령 행뿐**이다(측정: 20_개발\…\01_설계\_도구\related_law_summary.csv).
  법령DB에는 그 조문 의무가 있다 → 데모에 없는 것을 후보로 뽑고, 사람이 O/X 로 볼 표본을 만든다.

하는 일 / 하지 않는 일
  · 적용하지 않는다. 데모 자료 판(ops_*)·법령DB(build\) 는 읽기만 한다.
  · 판정·수범주체는 **법령DB 값을 그대로 옮긴다**(obl_axes.duty_verdict = obligation 인 의무만).
    새로 판정하지 않는다. 「용인시 해당 신호」는 판정이 아니라 **후보 선별 표시**다 —
    법령DB 수범주체 칸(unit_role.duty_subject_raw)을 먼저 보고, 비어 있으면 원문 주어 낱말을 본다.
    「낮음」도 지우지 않고 전체 후보 시트에 남긴다(원칙 4).
  · 데모 칸 제안값(code36·관리대상·이행 유형·주기·증빙)은 **제안**이다. 확신이 없으면 「확인 필요」.
      code36   ① 교육 자료 체크리스트 정답(mgmt_v1.0 checklist_gold) → ② 관리층 obl_category(중처법조문·체크리스트·규칙 R##)
               → ③ 기본값/없음 = 「확인 필요」(참고로 데모와 같은 추론값을 옆 칸에 적는다)
      관리대상 데모 duty_class 에서 그 법 행이 가장 많이 쓴 target_code(데이터에서) → 없으면 아래 LAW_TG 표
      이행 유형 build_yongin_3axis.py 와 같은 낱말 규칙(조 제목·의무명) → 안 걸리면 code36 기본 유형(추론 표시)
      주기     원문(앵커 조항호목 + 그 아래 호·목)에서 주기 낱말을 찾아 **원문 인용**으로 적는다. 없으면 「확인 필요」
      증빙     데모 duty_class 에서 같은 code36 이 가장 많이 쓴 evidence_kind
  · 입력은 발행 판만(_RELEASE.json). 데모 자료는 앱과 같은 규칙(ops_* 이름 역순, 처음 있는 판)으로 고른다.

입력
  법령DB  ..\..\build\_RELEASE.json → law · doc · unit · obl_master · obl_axes · unit_role
  데모    ..\ops_v*\seed\duty_class.csv(가장 새 판) · ..\mgmt_v1.0_20260920\(obl_category · checklist_gold · law_mgmt)
  서울시 붙임 1-1  20_개발\_데모_용인시_20260920\01_설계\_도구\seoul_annex11_articles.csv (표본 우선순위·표시용)
  명칭    30_데이터\_수집작업\_체크리스트_시설별\duty_master.py (의무조항 36 표준 명칭)

출력 (..\_검토\)
  (후보)데모의무_법률조문편입_20260921_v0.1.csv
  (검수)데모의무_법률조문편입_표본_20260921_v0.1.xlsx — 요약 · 표본 · 전체 후보 · 제안 규칙
"""
import collections
import csv
import io
import json
import os
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

csv.field_size_limit(2 ** 31 - 1)
if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = pathlib.Path(__file__).resolve().parent          # _build
DEMO = HERE.parent                                       # _데모_용인시_20260920
DB = DEMO.parent                                         # ADOMS_DB_v1
BUILD = DB / "build"
ROOT = DB.parents[2]                                     # 7.ADOMS 구현
CHK = DB.parent / "_체크리스트_시설별"
ANNEX = ROOT / "20_개발" / "_데모_용인시_20260920" / "01_설계" / "_도구" / "seoul_annex11_articles.csv"
MGMT = DEMO / "mgmt_v1.0_20260920"
OUT_DIR = DEMO / "_검토"
OUT_CSV = OUT_DIR / "(후보)데모의무_법률조문편입_20260921_v0.1.csv"
OUT_XLSX = OUT_DIR / "(검수)데모의무_법률조문편입_표본_20260921_v0.1.xlsx"

sys.path.insert(0, str(CHK))
sys.path.insert(0, str(DB / "script"))
from duty_master import MASTER, VERSION as DM_VERSION        # noqa: E402  표준 명칭 한 곳
from lib_naming import SUBITEM_HANGUL, SUBITEM_ORDER          # noqa: E402

NAME36 = {m[1]: m[2] for m in MASTER}
FONT = "맑은 고딕"

# ── 대상 법 (과업 지정 18법 + 도로 규칙 2 — 도로 규칙은 도로법 가족 문서로 정본에 있다: R-20260914-10) ──
TARGET_LAWS = [
    ("LAW-KR-000011", "시설물안전법"), ("LAW-KR-000006", "재난안전법"), ("LAW-KR-000162", "건축물관리법"),
    ("LAW-KR-000009", "소방시설법"), ("LAW-KR-000161", "화재예방법"), ("LAW-KR-000008", "전기안전관리법"),
    ("LAW-KR-000007", "승강기법"), ("LAW-KR-000114", "도로법"), ("LAW-KR-000138", "하천법"),
    ("LAW-KR-000174", "수도법"), ("LAW-KR-000175", "하수도법"), ("LAW-KR-000168", "도시철도법"),
    ("LAW-KR-000021", "철도안전법"), ("LAW-KR-000173", "지하안전법"), ("LAW-KR-000178", "기계설비법"),
    ("LAW-KR-000163", "실내공기질법"), ("LAW-KR-000166", "어린이놀이시설법"), ("LAW-KR-000119", "체육시설법"),
]
LAW_ABBR = dict(TARGET_LAWS)
STATUTORY = {"act": "법률", "presidential_decree": "대통령령(시행령)", "ministerial_ordinance": "부령(시행규칙)"}
LAYER_ORDER = {"act": 0, "presidential_decree": 1, "ministerial_ordinance": 2}

# 데모에 그 법 행이 없을 때만 쓰는 관리대상 기본값 (코드·이름은 데모 duty_class 의 TG 체계 그대로)
#   값: (TG, 확신) — 「확인 필요」면 제안 칸에 코드 대신 확인 필요를 쓴다
LAW_TG = {
    "LAW-KR-000162": ("TG13", "보통"),   # 건축물관리법 → 업무시설·복합건축물(데모 규칙 「건축물」 낱말)
    "LAW-KR-000009": ("TG20", "보통"),   # 소방시설법 → 소방대상물·소방시설
    "LAW-KR-000161": ("TG20", "보통"),   # 화재예방법 → 소방대상물·소방시설
    "LAW-KR-000173": ("TG33", "확인 필요"),  # 지하안전법 → 지하시설물(상·하수관·공동구·지하차도 …) — 한 TG 로 못 묶음
    "LAW-KR-000178": ("TG13", "확인 필요"),  # 기계설비법 → 건축물 기계설비 — 데모에 기계설비 TG 없음
    "LAW-KR-000163": ("TG24", "확인 필요"),  # 실내공기질법 → 다중이용시설 여러 종 — 한 TG 로 못 묶음
}

# ── 이행 유형 (build_yongin_3axis.py TYPES 와 같은 규칙 · 순서 그대로) ──
TYPES = [
    ("T03", "교육·훈련", r"교육|훈련|강습|연수|이수"),
    ("T02", "인력·자격·선임", r"선임|배치|자격|면허|기술인력|안전관리자|책임자\s*지정|담당자\s*지정|운영관리사|자위소방대|"
                        r"편성|근무|파견|복무|당직|관제요원|승무|인력|복장|안전장구"),
    ("T09", "비상대응·사고 처리", r"비상|응급|긴급|대피|피난|재난|사고|재해\s*발생|재발|구조|구급|복구|경보|화재\s*진압|"
                           r"상황\s*전파|대처|협조|수습|경계"),
    ("T08", "위생·보건", r"건강진단|건강검사|보건조치|감염|소독|예방접종|위생|방역"),
    ("T07", "물질·제품 관리", r"유해화학물질|독성가스|위험물|마약|농약|비료|식품|의약품|제조|수입|판매|유통|회수|폐기|품질|취급"),
    ("T04", "점검·검사·진단·측정", r"점검|검사|진단|측정|계측|시험|검정|심사|안전성\s*평가|실태\s*조사|확인|조사분석|감시|주요\s*부분"),
    ("T05", "시설·설비·구조 기준", r"설치|구조|성능|기준|한계|규격|치수|유량|사용수량|적합|유지관리|유지ㆍ관리|보수|보강|정비|"
                          r"설계|시공|감리|내화|방화|부적합설비"),
    ("T06", "작업·운행 안전조치", r"작업|운행|운전|조작|승차|적재|출입|금지|보호구|안전장치|방호|안전조치|준수사항|주의사항|전호"),
    ("T10", "기록·보고·게시", r"기록|보존|서류|제출|보고|신고|통보|게시|표지|표시|공시|공표|비치|작성|고시|공고|지정|등록|"
                       r"말소|협의|수리|신청|대장|서식|발송"),
    ("T01", "체계·규정·계획", r"계획|규정|매뉴얼|지침|위원회|협의체|체계|의견|절차|조직|운영|역할|예산|재원|보험|권고"),
]
TYPES = [(c, n, re.compile(p)) for c, n, p in TYPES]
TNAME = {c: n for c, n, _ in TYPES}
CODE2TYPE = {
    "I01": "T01", "I02": "T01", "I03": "T04", "I04": "T01", "I05": "T02", "I06": "T02", "I07": "T01",
    "I08": "T09", "I09": "T01", "I10": "T09", "I11": "T10", "I12": "T05", "I13": "T03", "I14": "T01",
    "F01": "T02", "F02": "T01", "F03": "T04", "F04": "T01", "F05": "T04", "F06": "T09", "F07": "T01",
    "F08": "T01", "F09": "T09", "F10": "T10", "F11": "T05", "F12": "T03", "F13": "T01",
    "M01": "T02", "M02": "T01", "M03": "T07", "M04": "T01", "M05": "T09", "M06": "T09", "M07": "T10",
    "M08": "T05", "M09": "T03",
}
# 데모가 기본값(관계법령 의무이행) 행을 이행 유형으로 다시 배정한 표 — 참고 추론값에만 쓴다
TYPE2CODE = {
    "F": {"T01": "F04", "T02": "F01", "T03": "F12", "T04": "F03", "T05": "F11", "T06": "F11",
          "T07": "F11", "T08": "F11", "T09": "F07", "T10": "F11"},
    "I": {"T01": "I12", "T02": "I06", "T03": "I13", "T04": "I12", "T05": "I12", "T06": "I12",
          "T07": "I12", "T08": "I12", "T09": "I08", "T10": "I12"},
    "M": {"T01": "M04", "T02": "M01", "T03": "M09", "T04": "M05", "T05": "M03", "T06": "M03",
          "T07": "M03", "T08": "M03", "T09": "M04", "T10": "M08"},
}

# ── 용인시 해당 신호 (판정 아님 · 후보 선별 표시) ─────────────────────────
#    판정 규칙 v1.3 §2 의 네 자격을 낱말로 가리키는 표시일 뿐이다. 글자로 거르지 않는다 → 「낮음」도 남긴다.
INDIV = re.compile(r"안전관리자|안전관리보조자|소방안전관리자|종사자|운전자|운전업무|관제업무|작업책임자|기술자|기관사|"
                   r"승무원|누구든지|이용자|승객|여객|상황근무자|최초로\s*접수한\s*자|안전관리기술자")
YES = re.compile(r"관리주체|관리청|재난관리책임기관|지방자치단체|시장ㆍ군수|시장·군수|군수|구청장|관계인|소유자|관리자|점유자|"
                 r"설치자|운영자|수도사업자|하수도관리|지하시설물관리자|발주한\s*자|발주자|선임한\s*자|관리기관|"
                 r"국가핵심기반|다중이용시설|시설의\s*장|공공기관")
UNCERTAIN = re.compile(r"지하개발사업자|사업시행자|개발사업자|위험물취급자|사용자|건축주|공사시행자|사업자|사업주체|수급인|"
                       r"시ㆍ도지사|시·도지사|특별시장|광역시장|도지사")
LOW = re.compile(r"장관|청장|처장|국토교통부|행정안전부|소방청|해양경찰|소방본부장|소방서장|제조|수입업자|판매|대여|중개|시공자|"
                 r"정비조직|점용자|운송사업자|관리업자|체육시설업자|검사기관|전문기관|진단기관|교육기관|협회|공단|공사업자|감리|"
                 r"지정을\s*받으려는|신청하려는|인증을\s*받으려는|등록하려는|사무취급기관|보험회사|보험사업자|교육훈련기관|"
                 r"형식승인을\s*받|성능인증을\s*받|안전인증을\s*받|기술원|공단|탐색구조본부|해양|평가단장|통제단장|현장지휘관|긴급구조기관")
TYPE_KO = {"business_owner": "사업주·관리주체", "chief_officer": "기관의 장", "safety_manager": "안전관리자(개인)",
           "other": "기타", "general_public": "일반인"}
SECTOR_KO = {"all": "공공·민간 모두", "central_gov": "중앙행정기관", "public_inst": "공공기관", "local_gov": "지방자치단체"}
ADMIN = re.compile(r"^\s*(?:제\d+조(?:의\d+)?\([^)]*\)\s*)?(?:[①-⑳]\s*)?(?:\d+\.\s*)?(?:국가|지방자치단체|시ㆍ도지사|시ㆍ도|"
                   r"시장ㆍ군수ㆍ구청장|특별자치시장|시장|군수|구청장|도지사|특별시장|광역시장|행정안전부장관|국토교통부장관|"
                   r"환경부장관|기후에너지환경부장관|고용노동부장관|소방청장|소방본부장|소방서장|관계\s*중앙행정기관의\s*장|"
                   r"중앙행정기관의\s*장|[가-힣]+부장관|[가-힣]+청장|관리청|도로관리청|하천관리청)"
                   r"(?:[ㆍ·,\s]+[가-힣ㆍ·\s]*)?(?:은|는|이|가)\s")
# R14(인력 선임·배치)가 진짜 인력 낱말에 걸렸는지 — 「안전관리체계」「안전관리규정」의 「안전관리」만 걸린 것은 인력이 아니다
R14_REAL = re.compile(r"선임|배치|책임관|관리자|관리인|담당자|기술인력|보호인력|자격|운영관리사|성실의무|자위소방대|편성")
POWER = re.compile(r"(허가|승인|지정|인가|명할|명하여|취소|정지|검사할|점검할|조사할|요청할|요구할|과태료|징수)")

# ── 주기 낱말 (원문 인용으로만 쓴다) ──
CYCLE = re.compile(
    r"(매년|매월|매분기|매반기|매주|매일|반기(?:마다|별로|에\s*1회)?|분기(?:마다|별로|에\s*1회)?|"
    r"연\s*\d+\s*회|년\s*\d+\s*회|월\s*\d+\s*회|"
    r"(?:\d+|한|두|세)\s*(?:년|개월|월|주|일)\s*(?:마다|에\s*(?:1|한)\s*(?:회|번)\s*이상|에\s*(?:1|한)\s*(?:회|번))|"
    r"(?:\d+)\s*(?:년|개월|일)\s*(?:이내|이상\s*보존|간\s*보존|동안\s*보존)|"
    r"지체\s*없이|즉시|수시로|정기적으로|정기(?:안전)?점검)")

OPS_DIRS = sorted([d for d in DEMO.iterdir() if d.is_dir() and d.name.startswith("ops_")], key=lambda p: p.name, reverse=True)


def rd(path, cols=None):
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            yield r if cols is None else {c: r.get(c, "") or "" for c in cols}


def newest_seed(table):
    for d in OPS_DIRS:
        p = d / "seed" / (table + ".csv")
        if p.exists():
            return d.name, p
    return None, None


def norm(s):
    return re.sub(r"[\s·ㆍ･‧・.]", "", s or "")


def path_key(p):
    """a10g2/p3/n1/mga → 정렬 튜플."""
    out = []
    for seg in (p or "").split("/"):
        m = re.fullmatch(r"([apn])(\d+)(?:g(\d+))?", seg)
        if m:
            out.append((m.group(1), int(m.group(2)), int(m.group(3) or 0)))
            continue
        m = re.fullmatch(r"m([a-z]+)", seg)
        if m and m.group(1) in SUBITEM_HANGUL:
            out.append(("m", SUBITEM_ORDER[SUBITEM_HANGUL[m.group(1)]], 0))
            continue
        out.append(("z", 999, 0))
    return tuple(out)


def art_label(no, br):
    return ("제%s조" % no) + (("의%s" % br) if br else "") if no else ""


def fit_of(text):
    """낱말 표시 → (등급, 걸린 낱말). 판정 아님."""
    t = text or ""
    if re.search(r"\(민간\)|민간관리주체|민간\s*관리주체", t):
        return "확인 필요", "민간 관리주체 한정"
    ind = INDIV.findall(t)
    t2 = INDIV.sub(" ", t)
    yes = YES.findall(t2)
    if yes:
        return "가능", "·".join(dict.fromkeys(yes))
    if ind:
        return "확인 필요(개인 의무)", "·".join(dict.fromkeys(ind))
    unc = UNCERTAIN.findall(t2)
    if unc:
        return "확인 필요", "·".join(dict.fromkeys(unc))
    low = LOW.findall(t2)
    if low:
        return "낮음", "·".join(dict.fromkeys(low))
    return "확인 필요", ""


SKIP_TAIL = re.compile(r"(에|에서|로|으로|경우|때|외|도|에게|까지|부터|마다|이하|이상|이내|중|간|내|전|후|상|한)$")


def subject_phrase(text):
    """원문 첫 문장에서 「…은/는」 주어 낱말(최대 5어절)을 찾는다. 없으면 ''."""
    s = re.sub(r"^\s*제\d+조(?:의\d+)?\([^)]*\)\s*", "", text or "")
    s = re.sub(r"^\s*[①-⑳]\s*", "", s)
    s = re.sub(r"^\s*\d+\.\s*", "", s)
    s = re.sub(r"<[^>]*>", "", s)
    s = re.split(r"하여야\s*한다|해야\s*한다|아니\s*된다|아니\s*한다", s)[0]
    words = s.split()
    for i, w in enumerate(words):
        m = re.fullmatch(r"(.+?)(은|는)[,]?", w)
        if not m:
            continue
        stem = m.group(1)
        if SKIP_TAIL.search(stem) or stem.endswith("에서") or len(stem) < 1:
            continue
        return " ".join(words[max(0, i - 4): i + 1])
    return ""


def cycle_quote(text):
    """주기 낱말 앞뒤 원문을 인용한다(최대 3곳)."""
    out = []
    for m in CYCLE.finditer(text or ""):
        a, b = max(0, m.start() - 18), min(len(text), m.end() + 18)
        q = re.sub(r"\s+", " ", text[a:b]).strip()
        out.append("「…%s…」" % q)
        if len(out) >= 3:
            break
    return " / ".join(out)


def main(overwrite=False):
    for p in (OUT_CSV, OUT_XLSX):
        if p.exists() and not overwrite:
            sys.exit("이미 있다 — 덮어쓰려면 --overwrite (사용자 검수 파일이면 덮어쓰지 않는다): %s" % p)
    OUT_DIR.mkdir(exist_ok=True)

    rel = json.load(open(BUILD / "_RELEASE.json", encoding="utf-8"))
    T = {k: BUILD / v["file"] for k, v in rel["tables"].items()}
    release = rel["release"]
    print("법령DB 발행 판", release, rel.get("state"))

    # ── 법령DB ─────────────────────────────────────────────
    law = {r["law_id"]: r for r in rd(T["law"], ["law_id", "title_ko", "abbr_ko", "status", "merged_into"])}
    tgt_law = set(LAW_ABBR)
    docs = {}
    for r in rd(T["doc"], ["doc_id", "law_id", "title_ko", "norm_form", "dup_status", "canonical_doc_id"]):
        lid = r["law_id"]
        if law.get(lid, {}).get("merged_into") in tgt_law:
            lid = law[lid]["merged_into"]
        if lid in tgt_law and r["norm_form"] in STATUTORY:
            r["fam"] = lid
            docs[r["doc_id"]] = r
    print("대상 문서(법률·시행령·시행규칙)", len(docs))

    units = {}
    by_art = collections.defaultdict(list)
    for r in rd(T["unit"], ["unit_id", "doc_id", "unit_path", "unit_label", "unit_type", "article_no",
                             "article_branch", "article_title", "unit_text", "is_deleted", "text_updated_by"]):
        if r["doc_id"] in docs:
            units[r["unit_id"]] = r
            by_art[(r["doc_id"], r["article_no"], r["article_branch"])].append(r)
    for k in by_art:
        by_art[k].sort(key=lambda u: path_key(u["unit_path"]))
    art_title = {}
    for k, us in by_art.items():
        for u in us:
            if u["article_title"]:
                art_title[k] = u["article_title"]
                break

    om = {r["obl_id"]: r for r in rd(T["obl_master"], ["obl_id", "anchor_unit_id", "title_ko", "obligation_group",
                                                        "cycle", "evidence_required", "nature", "review_status"])}
    ax = {}
    for r in rd(T["obl_axes"], ["obl_id", "anchor_unit_id", "doc_id", "role", "duty_verdict", "review_status"]):
        if r["doc_id"] in docs:
            ax[r["obl_id"]] = r
    ur = {}
    for r in rd(T["unit_role"], ["unit_id", "role", "duty_subject_raw", "duty_subject_type", "duty_subject_sector",
                                 "subject_src", "judgment_basis", "exclusion_reason", "judged_by"]):
        if r["unit_id"] in units:
            ur[r["unit_id"]] = r

    # ── 데모 · 관리층 ───────────────────────────────────────
    dc_ver, dc_path = newest_seed("duty_class")
    duty = list(rd(dc_path))
    demo_unit = {r["unit_id"] for r in duty if r["unit_id"]}
    demo_obl = {r["obl_id"] for r in duty if r["obl_id"]}
    demo_art = collections.defaultdict(list)
    for r in duty:
        u = units.get(r["unit_id"])
        if u:
            demo_art[(u["doc_id"], u["article_no"], u["article_branch"])].append(r["duty_key"])
    tg_name = {}
    law_tg = collections.defaultdict(collections.Counter)
    ev36 = collections.defaultdict(collections.Counter)
    for r in duty:
        if r["target_code"]:
            tg_name[r["target_code"]] = r["target_name"]
            law_tg[r["law_id"]][r["target_code"]] += 1
        for c in (r["code36"] or "").split(";"):
            if c and r["evidence_kind"]:
                ev36[c][r["evidence_kind"]] += 1
    print("데모 의무 자료", dc_ver, len(duty), "행")

    cat = collections.defaultdict(dict)       # obl_id → {area: row}
    for r in rd(MGMT / "obl_category.csv"):
        cat[r["obl_id"]][r["area"]] = r
    gold = {}
    for r in rd(MGMT / "checklist_gold.csv"):
        if r.get("3단 obl_id"):
            gold.setdefault(r["3단 obl_id"], (r["1단 코드"], r["2단 세부 이행업무(체크리스트)"]))
    lawmg = {r["law_id"]: r for r in rd(MGMT / "law_mgmt.csv")}

    annex = collections.defaultdict(list)     # (norm 문서명, 조, 가지) → 붙임 내용
    for r in rd(ANNEX):
        m = re.fullmatch(r"(\d+)(?:의(\d+))?", (r["article"] or "").strip())
        if m:
            annex[(norm(r["law_name"]), m.group(1), m.group(2) or "")].append(r["content"])

    # ── 후보 ───────────────────────────────────────────────
    rows = []
    stat_all = collections.Counter()
    for oid, a in ax.items():
        if a["duty_verdict"] != "obligation":
            stat_all[("판정 obligation 아님", docs[a["doc_id"]]["fam"])] += 1
            continue
        o = om.get(oid, {})
        uid = a["anchor_unit_id"]
        u = units.get(uid)
        if not u:
            continue
        d = docs[u["doc_id"]]
        fam = d["fam"]
        stat_all[("판정 obligation", fam)] += 1
        if uid in demo_unit or oid in demo_obl:
            stat_all[("데모에 이미 있음", fam)] += 1
            continue
        akey = (u["doc_id"], u["article_no"], u["article_branch"])
        art = by_art.get(akey, [])
        title = art_title.get(akey, "")
        own = u["unit_text"]
        sub_units = [x for x in art if x["unit_path"].startswith(u["unit_path"] + "/")]
        own_full = "\n".join([own] + [x["unit_text"] for x in sub_units])
        art_full = "\n".join(x["unit_text"] for x in art if x["unit_text"])
        role = ur.get(uid, {})

        # 용인시 해당 신호 — 법령DB 수범주체 칸 먼저, 비면 원문 주어 낱말
        raw = role.get("duty_subject_raw", "")
        if raw:
            fit, hit = fit_of(raw)
            fit_src = "법령DB 수범주체 「%s」" % raw
        else:
            ph = subject_phrase(own) or subject_phrase(art[0]["unit_text"] if art else "")
            fit, hit = fit_of(ph) if ph else ("확인 필요", "")
            fit_src = ("수범주체 공란 · 원문 주어 「%s」" % ph) if ph else "수범주체 공란 · 원문 주어 없음(시설·구조 기준 문장 등)"
        flags = []
        if ADMIN.search(own) and POWER.search(own):
            flags.append("행정청 권한 의심(판정규칙 §2 ③)")
        if "보험" in own_full:
            flags.append("보험 조항 — 판정규칙 §6-2(이행 증빙 층) 해당 확인")
        if "노력하여야" in own:
            flags.append("노력의무")
        if u["text_updated_by"]:
            flags.append("판정 뒤 본문 갱신(%s) — 재검토" % u["text_updated_by"])
        if u["is_deleted"] == "true":
            flags.append("삭제된 조항호목")
        if re.search(r"건설규칙|구조ㆍ시설\s*기준", d["title_ko"]):
            flags.append("설계·건설 기준 — 운영 중 시설 적용 여부 확인")
        if "노면전차" in d["title_ko"]:
            flags.append("노면전차 규칙 — 용인경전철이 노면전차인지 확인")
        if "채권" in d["title_ko"]:
            flags.append("채권 사무 — 안전 의무인지 확인")
        if demo_art.get(akey):
            flags.append("같은 조의 다른 조항호목이 데모에 있음(%d행)" % len(demo_art[akey]))

        # code36
        c_by = cat.get(oid, {})
        area = "F" if "F" in c_by else ("I" if "I" in c_by else ("M" if "M" in c_by else "F"))
        other = " · ".join("%s %s" % (k, v["code"]) for k, v in sorted(c_by.items()) if k != area)
        typ, typ_src = None, ""
        for c, n, rx in TYPES:
            if rx.search("%s %s" % (title, o.get("title_ko", ""))):
                typ, typ_src = c, "규칙(조 제목·의무명 낱말 「%s」)" % rx.search("%s %s" % (title, o.get("title_ko", ""))).group(0)
                break
        if oid in gold:
            code, conf = gold[oid][0], "높음"
            basis = "교육 자료 체크리스트 정답 — 「%s」" % gold[oid][1]
            hint = ""
        elif area in c_by and c_by[area]["basis_type"] in ("중처법조문", "체크리스트"):
            code, conf = c_by[area]["code"], "높음"
            basis = "관리층 obl_category(%s · %s)" % (c_by[area]["basis_type"], c_by[area]["basis_detail"])
            hint = ""
        elif area in c_by and c_by[area]["basis_type"] == "규칙" and not (
                c_by[area]["basis_detail"].startswith("R14") and not R14_REAL.search("%s %s" % (title, o.get("title_ko", "")))):
            code, conf = c_by[area]["code"], "보통"
            basis = "관리층 obl_category 낱말 규칙 %s — 조 제목·의무명에 걸림" % c_by[area]["basis_detail"]
            hint = ""
        else:
            t0 = typ or CODE2TYPE.get("F11")
            if area in c_by and c_by[area]["basis_type"] == "규칙":   # R14 가 「안전관리」 낱말에만 걸린 경우
                t0 = typ or "T05"
            hint = TYPE2CODE[area][t0]
            code, conf = "확인 필요", "낮음"
            basis = ("관리층 obl_category 기본값(%s 관계법령 의무이행) — 규칙에 안 걸림" % c_by[area]["code"]
                     if area in c_by and c_by[area]["basis_type"] != "규칙" else
                     "관리층 obl_category %s(%s)가 R14 인데 「안전관리」 낱말(안전관리체계·안전관리규정 등)에만 걸림 — 인력 배치가 아님" % (
                         c_by[area]["code"], c_by[area]["basis_detail"]) if area in c_by else "관리층 obl_category 에 없음(발행 판 R-20260917-02 뒤 생긴 의무이거나 영역 미배정)")
            basis += " · 데모식 추론값 %s(%s)" % (hint, NAME36.get(hint, ""))
        code_for_type = code if code in NAME36 else hint
        if not typ:
            typ = CODE2TYPE.get(code_for_type, "T05")
            typ_src = "추론 — 낱말 규칙에 안 걸려 %s 기본 유형" % code_for_type
        # 관리대상
        cnt = law_tg.get(fam)
        if cnt:
            tg, n = cnt.most_common(1)[0]
            tot = sum(cnt.values())
            tg_conf = "보통" if n / tot >= 0.6 else "확인 필요"
            tg_basis = "데모에서 이 법 행 %d 중 %d가 %s" % (tot, n, tg)
        elif fam in LAW_TG:
            tg, tg_conf = LAW_TG[fam]
            tg_basis = "데모에 이 법 행 없음 — 법 기본값"
        else:
            tg, tg_conf, tg_basis = "", "확인 필요", "기본값 없음"
        if raw and re.search(r"재난관리책임기관", raw):
            tg, tg_conf, tg_basis = "TG24", "보통", "수범주체가 재난관리책임기관 — 데모 관례 TG24(기관 전체)"
        # 주기
        cq = cycle_quote(own_full)
        lawdb_cycle = o.get("cycle", "")
        if cq:
            cyc = cq
        else:
            cyc = "확인 필요(원문에 주기 낱말 없음 — 하위법령·별표 확인)"
        ev = ev36[code_for_type].most_common(1)[0][0] if ev36.get(code_for_type) else "확인 필요"
        if code not in NAME36:
            ev = "확인 필요(code36 미확정 · 추론값 기준: %s)" % ev
        an = annex.get((norm(d["title_ko"]), u["article_no"], u["article_branch"]))
        lg = lawmg.get(fam, {})
        rows.append({
            "cand_id": "",
            "law_abbr": LAW_ABBR[fam], "law_id": fam, "law_title": law[fam]["title_ko"],
            "doc_id": d["doc_id"], "doc_title": d["title_ko"], "layer": STATUTORY[d["norm_form"]],
            "unit_id": uid, "unit_path": u["unit_path"], "unit_label": u["unit_label"],
            "article": art_label(u["article_no"], u["article_branch"]), "article_title": title,
            "obl_id": oid, "obl_title": o.get("title_ko", ""),
            "lawdb_verdict": a["duty_verdict"], "lawdb_role": a["role"], "lawdb_review": a["review_status"],
            "lawdb_subject_raw": raw, "lawdb_subject_type": TYPE_KO.get(role.get("duty_subject_type", ""), role.get("duty_subject_type", "")),
            "lawdb_subject_sector": SECTOR_KO.get(role.get("duty_subject_sector", ""), role.get("duty_subject_sector", "")),
            "lawdb_subject_src": role.get("subject_src", ""),
            "yongin_fit": fit, "yongin_fit_hit": hit, "yongin_fit_basis": fit_src,
            "is_candidate": "N" if fit == "낮음" else "Y",
            "flags": " · ".join(flags),
            "seoul_annex": ("Y — " + " / ".join(an)) if an else "",
            "p_area": area, "p_code36": code, "p_code36_name": NAME36.get(code, ""), "p_code36_conf": conf,
            "p_code36_basis": basis, "p_code36_hint": hint, "p_code36_other_area": other,
            "p_target_code": tg if tg_conf != "확인 필요" else "확인 필요",
            "p_target_hint": tg, "p_target_name": tg_name.get(tg, ""), "p_target_basis": tg_basis,
            "p_impl_type": typ, "p_impl_type_name": TNAME[typ], "p_impl_basis": typ_src,
            "p_cycle_quote": cyc, "lawdb_cycle": lawdb_cycle,
            "p_evidence_kind": ev,
            "p_law_group": lg.get("law_family", ""), "p_law_group_name": lg.get("law_family_name", ""),
            "p_task_name": "%s %s(%s)" % (d["title_ko"], art_label(u["article_no"], u["article_branch"]), title),
            "anchor_text": own_full, "article_text": art_full,
        })
    rows.sort(key=lambda r: ([k for k, _ in TARGET_LAWS].index(r["law_id"]), LAYER_ORDER[
        {v: k for k, v in STATUTORY.items()}[r["layer"]]], path_key(r["unit_path"]), r["doc_id"]))
    for i, r in enumerate(rows, 1):
        r["cand_id"] = "CND-%04d" % i

    # ── 표본 ───────────────────────────────────────────────
    fit_rank = {"가능": 0, "확인 필요": 1, "확인 필요(개인 의무)": 2, "낮음": 3}
    by_law = collections.defaultdict(list)
    for r in rows:
        if r["is_candidate"] == "Y":
            by_law[r["law_id"]].append(r)
    sample = []
    for lid, _ in TARGET_LAWS:
        cands = by_law.get(lid, [])
        n = 3 if len(cands) < 30 else (4 if len(cands) < 100 else 5)
        pri = sorted(cands, key=lambda r: (0 if r["seoul_annex"] else 1, fit_rank[r["yongin_fit"]],
                                           0 if r["p_code36_conf"] != "낮음" else 1, r["cand_id"]))
        picked, seen_art, seen_layer = [], set(), set()
        # 1차: 붙임 조 · 가능 — 조 겹치지 않게
        for r in pri:
            if len(picked) >= n:
                break
            k = (r["doc_id"], r["article"])
            if k in seen_art:
                continue
            if r["seoul_annex"] and r["yongin_fit"] == "가능":
                picked.append(r); seen_art.add(k); seen_layer.add(r["layer"])
        # 2차: 아직 없는 층(시행령·시행규칙) · 도로 규칙
        for r in pri:
            if len(picked) >= n:
                break
            k = (r["doc_id"], r["article"])
            if k in seen_art:
                continue
            need_road = lid == "LAW-KR-000114" and "규칙" in r["doc_title"] and "도로법" not in r["doc_title"] \
                and not any("규칙" in x["doc_title"] and "도로법" not in x["doc_title"] for x in picked)
            if r["layer"] not in seen_layer or need_road:
                picked.append(r); seen_art.add(k); seen_layer.add(r["layer"])
        # 3차: 「확인 필요」 한 행은 꼭 넣는다(사람이 경계를 보게)
        if len(picked) < n and not any(x["yongin_fit"] != "가능" for x in picked):
            for r in pri:
                k = (r["doc_id"], r["article"])
                if r["yongin_fit"] != "가능" and k not in seen_art:
                    picked.append(r); seen_art.add(k); break
        for r in pri:
            if len(picked) >= n:
                break
            k = (r["doc_id"], r["article"])
            if k not in seen_art:
                picked.append(r); seen_art.add(k)
        sample.extend(picked)

    # ── CSV ────────────────────────────────────────────────
    cols = list(rows[0].keys())
    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    # ── 엑셀 ───────────────────────────────────────────────
    wb = Workbook()
    hfont = Font(name=FONT, bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor="305496")
    bfont = Font(name=FONT, size=10)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {"가능": "E2EFDA", "확인 필요": "FFF2CC", "확인 필요(개인 의무)": "FCE4D6", "낮음": "EDEDED"}
    rev_fill = PatternFill("solid", fgColor="DDEBF7")

    def sheet(ws, head, data, widths, fit_col=None, review_cols=()):
        ws.append(head)
        for c in ws[1]:
            c.font, c.fill, c.alignment, c.border = hfont, hfill, Alignment(wrap_text=True, vertical="center"), border
        for row in data:
            ws.append([("" if v is None else (str(v)[:32000] if isinstance(v, str) else v)) for v in row])
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.font, c.alignment, c.border = bfont, wrap, border
            if fit_col is not None:
                v = r[fit_col].value
                if v in fills:
                    r[fit_col].fill = PatternFill("solid", fgColor=fills[v])
            for rc in review_cols:
                r[rc].fill = rev_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # 요약
    ws = wb.active
    ws.title = "요약"
    cand = [r for r in rows if r["is_candidate"] == "Y"]
    S = []
    S.append(["데모 의무 — 법률·시행령·시행규칙 조문 의무 편입 후보 (적용 전 · 검수용)", "", "", "", "", "", "", ""])
    S.append(["법령DB 발행 판 %s · 데모 의무 자료 %s duty_class %d행 · 관리층 mgmt_v1.0 · 의무조항 36 표준 명칭 %s · 만든 코드 _build\\propose_statute_duties_v07.py"
              % (release, dc_ver, len(duty), DM_VERSION), "", "", "", "", "", "", ""])
    S.append(["", "", "", "", "", "", "", ""])
    S.append(["법", "법령DB 판정 obligation (법률·시행령·시행규칙)", "데모에 이미 있음(unit_id·obl_id)", "데모에 없음",
              "후보(가능)", "후보(확인 필요)", "후보 아님(낮음 — 지우지 않음)", "표본"])
    smp_n = collections.Counter(r["law_id"] for r in sample)
    tot = collections.Counter()
    for lid, ab in TARGET_LAWS:
        rs = [r for r in rows if r["law_id"] == lid]
        a1 = stat_all[("판정 obligation", lid)]
        a2 = stat_all[("데모에 이미 있음", lid)]
        c_yes = sum(1 for r in rs if r["yongin_fit"] == "가능")
        c_chk = sum(1 for r in rs if r["yongin_fit"].startswith("확인 필요"))
        c_low = sum(1 for r in rs if r["yongin_fit"] == "낮음")
        S.append(["%s(%s)" % (ab, law[lid]["title_ko"]), a1, a2, len(rs), c_yes, c_chk, c_low, smp_n[lid]])
        for k, v in zip(("a1", "a2", "n", "y", "c", "l", "s"), (a1, a2, len(rs), c_yes, c_chk, c_low, smp_n[lid])):
            tot[k] += v
    S.append(["합계", tot["a1"], tot["a2"], tot["n"], tot["y"], tot["c"], tot["l"], tot["s"]])
    S.append(["", "", "", "", "", "", "", ""])
    lay = collections.Counter((r["layer"], r["is_candidate"]) for r in rows)
    S.append(["층별(후보 Y / 낮음 N)"] + ["%s  Y %d · N %d" % (l, lay[(l, "Y")], lay[(l, "N")]) for l in STATUTORY.values()] + ["", "", "", ""])
    cc = collections.Counter(r["p_code36_conf"] for r in cand)
    S.append(["후보의 code36 제안 확신", "높음 %d" % cc["높음"], "보통 %d" % cc["보통"], "낮음(확인 필요) %d" % cc["낮음"], "", "", "", ""])
    fl = collections.Counter()
    for r in cand:
        for f_ in r["flags"].split(" · "):
            if f_:
                fl[re.sub(r"\(\d+행\)|\(R-[^)]*\)", "", f_)] += 1
    S.append(["후보에 붙은 주의 표시", "", "", "", "", "", "", ""])
    for k, v in fl.most_common():
        S.append(["", k, v, "", "", "", "", ""])
    cyc_hit = sum(1 for r in cand if not r["p_cycle_quote"].startswith("확인 필요"))
    S.append(["주기 원문 인용이 잡힌 후보", "%d / %d" % (cyc_hit, len(cand)), "", "", "", "", "", ""])
    S.append(["서울시 붙임 1-1 조에 해당하는 후보", sum(1 for r in cand if r["seoul_annex"]), "", "", "", "", "", ""])
    S.append(["", "", "", "", "", "", "", ""])
    S.append(["읽는 법", "", "", "", "", "", "", ""])
    for t in [
        "판정·수범주체 칸은 법령DB 값을 그대로 옮겼다. 이 표는 새로 판정하지 않는다.",
        "「용인시 해당 신호」는 판정이 아니라 선별 표시다 — 법령DB 수범주체 칸 → 비어 있으면 원문 주어 낱말. 「낮음」(중앙행정기관·제조업자·검사기관 등)도 전체 후보 시트에 남겼다.",
        "제안 값은 모두 검수 전이다. 확신이 없으면 「확인 필요」로 두고 참고 추론값을 옆 칸에 적었다.",
        "표본 시트의 파란 칸(편입·code36·관리대상·이행 유형·주기·증빙)에 O/X 를 적는다. 표본 결과로 규칙을 고친 뒤에만 전체 후보로 넓힌다.",
    ]:
        S.append(["", t, "", "", "", "", "", ""])
    for row in S:
        ws.append(row)
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 20
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if c.row != 1:
                c.font = Font(name=FONT, size=10, bold=(c.row == 4))
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # 표본
    head = ["표본", "cand_id", "법", "층", "조문", "조 제목", "법령DB 판정(역할)", "법령DB 수범주체(유형·부문)",
            "용인시 해당 신호", "신호 근거", "주의 표시", "서울 붙임 1-1", "앵커 원문(조항호목 + 아래 호·목)", "조 전문",
            "제안 code36", "code36 근거", "제안 관리대상", "관리대상 근거", "제안 이행 유형", "이행 유형 근거",
            "주기(원문 인용)", "법령DB 주기 칸", "제안 증빙 종류", "obl_id", "unit_id",
            "편입 O/X", "code36 O/X", "관리대상 O/X", "이행 유형 O/X", "주기 O/X", "증빙 O/X", "메모"]
    data = []
    for i, r in enumerate(sample, 1):
        data.append([i, r["cand_id"], r["law_abbr"], r["layer"], "%s %s" % (r["doc_title"], r["unit_label"]), r["article_title"],
                     "%s (%s)" % (r["lawdb_verdict"], r["lawdb_role"]),
                     "%s / %s · %s" % (r["lawdb_subject_raw"] or "(공란)", r["lawdb_subject_type"] or "-", r["lawdb_subject_sector"] or "-"),
                     r["yongin_fit"], r["yongin_fit_basis"] + ((" → 걸린 낱말 " + r["yongin_fit_hit"]) if r["yongin_fit_hit"] else ""),
                     r["flags"], r["seoul_annex"], r["anchor_text"], r["article_text"],
                     ("%s %s" % (r["p_code36"], r["p_code36_name"])).strip(), "[%s] %s" % (r["p_code36_conf"], r["p_code36_basis"]),
                     ("%s %s" % (r["p_target_code"], r["p_target_name"] if r["p_target_code"] != "확인 필요" else "(참고 %s %s)" % (r["p_target_hint"], r["p_target_name"]))).strip(),
                     r["p_target_basis"], "%s %s" % (r["p_impl_type"], r["p_impl_type_name"]), r["p_impl_basis"],
                     r["p_cycle_quote"], r["lawdb_cycle"], r["p_evidence_kind"], r["obl_id"], r["unit_id"],
                     "", "", "", "", "", "", ""])
    ws2 = wb.create_sheet("표본")
    sheet(ws2, head, data, [5, 10, 10, 12, 26, 18, 14, 22, 11, 30, 26, 18, 60, 70, 16, 34, 18, 26, 16, 26, 34, 10, 26, 13, 13,
                            8, 8, 8, 8, 8, 8, 24], fit_col=8, review_cols=range(25, 32))
    for r in ws2.iter_rows(min_row=2):
        ws2.row_dimensions[r[0].row].height = 180

    # 전체 후보
    ws3 = wb.create_sheet("전체 후보")
    head3 = ["cand_id", "후보", "용인시 해당 신호", "신호 근거", "법", "층", "문서", "조문", "조 제목", "의무명(법령DB)",
             "법령DB 판정", "역할", "법령DB 수범주체", "주의 표시", "서울 붙임", "제안 code36", "확신", "code36 근거",
             "참고 추론값", "제안 관리대상", "관리대상 근거", "제안 이행 유형", "주기(원문 인용)", "제안 증빙 종류",
             "앵커 원문", "obl_id", "unit_id"]
    data3 = [[r["cand_id"], r["is_candidate"], r["yongin_fit"], r["yongin_fit_basis"], r["law_abbr"], r["layer"], r["doc_title"],
              r["unit_label"], r["article_title"], r["obl_title"], r["lawdb_verdict"], r["lawdb_role"],
              r["lawdb_subject_raw"], r["flags"], "Y" if r["seoul_annex"] else "",
              ("%s %s" % (r["p_code36"], r["p_code36_name"])).strip(), r["p_code36_conf"], r["p_code36_basis"], r["p_code36_hint"],
              r["p_target_code"] if r["p_target_code"] != "확인 필요" else "확인 필요(참고 %s)" % r["p_target_hint"],
              r["p_target_basis"], "%s %s" % (r["p_impl_type"], r["p_impl_type_name"]), r["p_cycle_quote"], r["p_evidence_kind"],
              r["anchor_text"][:1500], r["obl_id"], r["unit_id"]] for r in rows]
    sheet(ws3, head3, data3, [10, 6, 11, 30, 10, 12, 22, 16, 18, 22, 10, 12, 18, 26, 6, 16, 7, 34, 8, 16, 26, 16, 30, 26, 60, 13, 13],
          fit_col=2)

    # 제안 규칙
    ws4 = wb.create_sheet("제안 규칙")
    rules = [
        ["칸", "어떻게 채웠나", "확인 필요가 되는 경우"],
        ["후보 범위", "법령DB 발행 판 obl_axes.duty_verdict = obligation · 문서가 대상 법 가족의 법률·시행령·시행규칙(부령 포함 · 도로 규칙 2 포함) · 데모 duty_class 에 같은 unit_id 나 obl_id 가 없는 것", "—"],
        ["용인시 해당 신호", "unit_role.duty_subject_raw(법령DB) 낱말 → 관리주체·관리청·재난관리책임기관·지방자치단체·시장·군수·소유자·관리자·점유자·설치자·운영자 등 = 가능 / 선임된 개인·종사자·운전자·이용자 = 확인 필요(개인 의무) / 사업시행자·사업자·시·도지사 등 = 확인 필요 / 장관·청장·제조·수입·판매업자·검사기관·지정 신청자 등 = 낮음. 법령DB 칸이 비면 원문 첫 「…은/는」 주어에 같은 낱말표를 쓴다", "주어가 없는 문장(시설·구조 기준) · 낱말표에 없는 주어"],
        ["code36", "① 교육 자료 체크리스트 정답(checklist_gold) → ② 관리층 obl_category 의 중처법조문·체크리스트 근거 = 높음 / 낱말 규칙 R## = 보통", "obl_category 가 기본값(관계법령 의무이행)이거나 행이 없을 때 → 「확인 필요」 + 데모식 추론값(이행 유형 → 36)"],
        ["관리대상", "데모 duty_class 에서 그 법 행이 가장 많이 쓴 target_code(60% 이상이면 보통) · 법령DB 수범주체가 재난관리책임기관이면 TG24(데모 관례)", "데모에 그 법 행이 없거나 한 TG 로 묶이지 않는 법(지하안전법·기계설비법·실내공기질법)"],
        ["이행 유형", "build_yongin_3axis.py 와 같은 낱말 규칙(조 제목 + 법령DB 의무명)", "규칙에 안 걸리면 code36 의 기본 유형(추론 표시)"],
        ["주기", "앵커 조항호목과 그 아래 호·목 원문에서 매년·반기·N년마다·지체 없이·즉시·N일 이내 등 낱말을 찾아 앞뒤 원문을 인용", "원문에 주기 낱말이 없을 때(대개 시행령·시행규칙·별표에 위임) — 법령DB 주기 칸은 참고로만 옆에 둔다"],
        ["증빙 종류", "데모 duty_class 에서 같은 code36 이 가장 많이 쓴 evidence_kind", "code36 이 확인 필요면 추론값 기준으로 적고 확인 필요 표시"],
        ["주의 표시", "행정청 권한 의심(주어가 행정청 + 허가·지정·명령 등 낱말) · 보험 조항(§6-2) · 노력의무 · 판정 뒤 본문 갱신 · 설계·건설 기준 · 노면전차 규칙 · 채권 사무 · 같은 조가 데모에 있음", "표시일 뿐 판정을 바꾸지 않는다"],
    ]
    sheet(ws4, rules[0], rules[1:], [16, 90, 60])

    wb.save(OUT_XLSX)
    print("후보 CSV", OUT_CSV.name, len(rows), "행 (후보 Y %d)" % len(cand))
    print("검수표", OUT_XLSX.name, "표본", len(sample))
    for lid, ab in TARGET_LAWS:
        rs = [r for r in rows if r["law_id"] == lid]
        print("  %-10s 데모에 없음 %4d · 가능 %4d · 확인 필요 %4d · 낮음 %4d · 표본 %d" % (
            ab, len(rs), sum(r["yongin_fit"] == "가능" for r in rs), sum(r["yongin_fit"].startswith("확인") for r in rs),
            sum(r["yongin_fit"] == "낮음" for r in rs), smp_n[lid]))


if __name__ == "__main__":
    main(overwrite="--overwrite" in sys.argv)
